#!/usr/bin/env python3
"""Export actual JSON/PPTX speaker text to Markdown, offline HTML, and optional DOCX.
This utility formats supplied text; it does not generate speech content with an LLM.
"""
from __future__ import annotations

import argparse
import html
import json
import math
import re
import sys
from pathlib import Path

CJK = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
WORDS = re.compile(r"[A-Za-z]+(?:[-'][A-Za-z]+)*")
NUMBERS = re.compile(r"\d+(?:[.,]\d+)*")


def clean_text(value):
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def numeric(value, field, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{field} must be a finite number")
    if value < 0 or (positive and value == 0):
        raise ValueError(f"{field} must be {'positive' if positive else 'non-negative'}")
    return float(value)


def text_list(value, field):
    if not isinstance(value, list) or any(not isinstance(item, str) for item in value):
        raise ValueError(f"{field} must be a list of strings")
    return list(value)


def normalized_slide(raw, index):
    if not isinstance(raw, dict):
        raise ValueError(f"Slide {index}: expected an object")
    sid = raw.get("id")
    if not isinstance(sid, (str, int)) or isinstance(sid, bool) or not str(sid).strip():
        raise ValueError(f"Slide {index}: id must be a non-empty string or integer")
    title = raw.get("title")
    notes = raw.get("notes", "")
    if not isinstance(title, str) or not isinstance(notes, str):
        raise ValueError(f"Slide {index}: title and notes must be strings")
    duration = raw.get("duration_seconds")
    if duration is not None:
        duration = numeric(duration, f"Slide {index}.duration_seconds")
    return {
        "id": str(sid), "title": clean_text(title), "notes": clean_text(notes),
        "cues": text_list(raw.get("cues", []), f"Slide {index}.cues"),
        "sources": text_list(raw.get("sources", []), f"Slide {index}.sources"),
        "duration_seconds": duration,
    }


def validate_deck(raw):
    if not isinstance(raw, dict) or not isinstance(raw.get("slides"), list) or not raw["slides"]:
        raise ValueError("Input must contain a non-empty slides list")
    if not isinstance(raw.get("title", ""), str):
        raise ValueError("Deck title must be a string")
    slides = [normalized_slide(row, idx) for idx, row in enumerate(raw["slides"], 1)]
    ids = [row["id"] for row in slides]
    if len(ids) != len(set(ids)):
        raise ValueError("Slide ids must be unique")
    return {"title": clean_text(raw.get("title", "")) or "演讲稿", "slides": slides}


def read_pptx(path):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ValueError("PPTX input/check requires python-pptx") from exc
    presentation = Presentation(str(path))
    slides = []
    for position, slide in enumerate(presentation.slides, 1):
        markers = [shape for shape in slide.shapes if shape.name.startswith("ppa-title:")]
        if len(markers) > 1:
            raise ValueError(f"PPTX slide {position} contains multiple ppa-title markers")
        if markers:
            marker = markers[0]
            sid = marker.name[len("ppa-title:"):]
            title = marker.text if marker.has_text_frame else ""
        else:
            sid = str(slide.slide_id)
            title_shape = slide.shapes.title
            title = title_shape.text if title_shape is not None else ""
            if not title:
                title = next((s.text for s in slide.shapes if s.has_text_frame and s.text.strip()), "")
        notes = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                notes = frame.text
        slides.append({"id": sid, "title": title, "notes": notes, "cues": [], "sources": []})
    title = presentation.core_properties.title or Path(path).stem
    return validate_deck({"title": title, "slides": slides})


def read_input(path):
    suffix = path.suffix.lower()
    if suffix == ".json":
        return validate_deck(json.loads(path.read_text(encoding="utf-8-sig")))
    if suffix == ".pptx":
        return read_pptx(path)
    raise ValueError("Input must be a .json storyboard or .pptx presentation")


def check_sync(expected, actual):
    issues = []
    a, b = expected["slides"], actual["slides"]
    if len(a) != len(b):
        issues.append(f"slide count differs: JSON {len(a)}, PPTX {len(b)}")
    for position, (left, right) in enumerate(zip(a, b), 1):
        for field in ("id", "title", "notes"):
            if left[field] != right[field]:
                detail = f" ({left[field]!r} vs {right[field]!r})" if field == "id" else ""
                issues.append(f"slide {position}: {field} differs{detail}")
    if issues:
        raise ValueError("JSON/PPTX synchronization failed; read the current deck and reconcile before export:\n- "
                         + "\n- ".join(issues[:20]))
    return {"status": "passed", "checked": ["page count", "current order", "id", "title", "actual notes"]}


def measure_notes(notes, zh_cpm, en_wpm, pause_seconds):
    chars = len(CJK.findall(notes))
    words = len(WORDS.findall(notes))
    number_tokens = len(NUMBERS.findall(notes))
    # Numbers count as an additional spoken token; complex formulas need actual rehearsal.
    seconds = math.ceil(60 * chars / zh_cpm + 60 * (words + number_tokens) / en_wpm
                        + (pause_seconds if notes.strip() else 0))
    return {"chinese_characters": chars, "english_words": words, "number_tokens": number_tokens,
            "estimated_seconds": seconds}


def mmss(seconds):
    total = max(0, int(round(seconds)))
    return f"{total // 60:02d}:{total % 60:02d}"


def prepare(decks, zh_cpm, en_wpm, pause_seconds):
    cumulative = 0
    missing, warnings = [], []
    for pos, slide in enumerate(decks["slides"], 1):
        timing = measure_notes(slide["notes"], zh_cpm, en_wpm, pause_seconds)
        chosen = slide["duration_seconds"]
        timing["basis"] = "storyboard_plan" if chosen is not None else "text_estimate"
        timing["scheduled_seconds"] = chosen if chosen is not None else timing["estimated_seconds"]
        timing["start_seconds"] = cumulative
        cumulative += timing["scheduled_seconds"]
        timing["end_seconds"] = cumulative
        slide["timing"] = timing
        if not slide["notes"]:
            missing.append({"position": pos, "id": slide["id"], "title": slide["title"]})
        if chosen is not None and chosen < timing["estimated_seconds"]:
            warnings.append(f"Slide {pos} ({slide['id']}): planned time is shorter than the text estimate; rehearse or trim.")
    decks["timing"] = {"zh_characters_per_minute": zh_cpm, "english_words_per_minute": en_wpm,
                      "pause_seconds_per_nonempty_slide": pause_seconds,
                      "scheduled_seconds": cumulative,
                      "estimated_seconds": sum(s["timing"]["estimated_seconds"] for s in decks["slides"]),
                      "label": "Text timing is an estimate, not a measured rehearsal. Planned durations come from the supplied storyboard."}
    return {"missing_notes": missing, "warnings": warnings}


def render_markdown(deck):
    lines = [f"# {deck['title']}", "",
             "本文件导出已有逐页口语稿；没有调用语言模型生成或补写缺失内容。", "",
             f"计划总时长：{mmss(deck['timing']['scheduled_seconds'])}；"
             f"文字估时：{mmss(deck['timing']['estimated_seconds'])}。估计不能代替实际试讲。", "",
             f"估算设置：中文 {deck['timing']['zh_characters_per_minute']:g} 字/分钟；"
             f"英文 {deck['timing']['english_words_per_minute']:g} 词/分钟；"
             f"每个有稿页面另加 {deck['timing']['pause_seconds_per_nonempty_slide']:g} 秒停顿。数字段近似按一词；公式/演示/互动需额外实测。", ""]
    for pos, slide in enumerate(deck["slides"], 1):
        t = slide["timing"]
        basis = "故事板计划" if t["basis"] == "storyboard_plan" else "文字估计"
        lines.extend([f"## {pos}. {slide['title'] or '未提供标题'}", "",
                      f"页面 ID：{slide['id']} · {mmss(t['start_seconds'])}–{mmss(t['end_seconds'])}"
                      f" · 本页 {mmss(t['scheduled_seconds'])}（{basis}）", "",
                      slide["notes"] or "【缺少原始备注／口语稿；未自动补写】", ""])
        if slide["cues"]:
            lines.extend(["**提示卡**", ""] + [f"- {cue}" for cue in slide["cues"]] + [""])
        if slide["sources"]:
            lines.extend(["**来源**", ""] + [f"- {source}" for source in slide["sources"]] + [""])
    return "\n".join(lines) + "\n"


HTML_TEMPLATE = """<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__ · 演讲稿</title><style>
:root{color-scheme:light;--ink:#172438;--muted:#5e6e83;--accent:#146779;--line:#dce5eb;--paper:#fff;--wash:#f2f5f7}
*{box-sizing:border-box}body{margin:0;background:var(--wash);color:var(--ink);font-family:"Microsoft YaHei","PingFang SC",Arial,sans-serif;line-height:1.65}
header{padding:24px clamp(18px,4vw,56px) 18px;background:var(--paper);border-bottom:1px solid var(--line)}
h1{font-size:clamp(22px,3vw,34px);line-height:1.25;margin:0 0 8px}header p{margin:4px 0;color:var(--muted);font-size:13px}
.toolbar{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-top:18px}
button,select{border:1px solid #b8cbd4;background:white;color:var(--ink);border-radius:8px;padding:9px 13px;font:inherit;font-size:14px;cursor:pointer;min-height:42px}
button:hover{background:#e9f3f5}button:focus-visible,select:focus-visible{outline:3px solid #31889a;outline-offset:2px}
button[aria-pressed="true"]{background:var(--accent);border-color:var(--accent);color:#fff}
button:disabled{opacity:.4;cursor:default}.clock{margin-left:auto;font-variant-numeric:tabular-nums;background:#edf4f6;padding:6px 12px;border-radius:8px}
main{max-width:1050px;margin:22px auto;padding:0 18px}.notice{color:#7e4318;background:#fff3e4;padding:12px 16px;border-radius:10px;margin-bottom:16px}
.slide{padding:clamp(22px,4vw,46px);background:white;border:1px solid var(--line);border-radius:16px;box-shadow:0 10px 35px #18233308}
.slide[hidden]{display:none}.slide h2{font-size:clamp(23px,3vw,34px);line-height:1.35;margin:0 0 12px}.meta{font-size:13px;color:var(--muted);margin-bottom:25px;font-variant-numeric:tabular-nums}
.notes{white-space:pre-wrap;overflow-wrap:anywhere;font-size:22px;line-height:1.9}.cues{font-size:25px;line-height:1.8;padding-left:1.2em}.empty{font-size:17px;color:#87511c}
details{margin-top:24px;border-top:1px solid var(--line);padding-top:14px;color:var(--muted);font-size:14px}details li{overflow-wrap:anywhere}
body[data-mode="cards"] .notes{display:none}body[data-mode="script"] .cues{display:none}
body[data-mode="all"] .slide{margin-bottom:20px}body[data-mode="all"] .cues{font-size:17px}
footer{text-align:center;color:var(--muted);font-size:13px;padding:18px}
@media(max-width:620px){.clock{margin-left:0;width:100%}.toolbar button{flex:1 1 100px;white-space:nowrap}.toolbar select{width:100%;order:-1}.notes{font-size:19px}.cues{font-size:22px}}
@media print{header .toolbar,footer,.notice{display:none}body{background:white}main{padding:0;margin:0;max-width:none}.slide,.slide[hidden]{display:block!important;break-before:page;box-shadow:none;border:0;border-radius:0}.notes{display:block!important}.cues{display:block!important;font-size:12pt}.slide:first-child{break-before:auto}details{display:block}.notes{font-size:12pt}header{padding:0 0 10px}}
</style></head><body data-mode="script">
<header><h1>__TITLE__</h1><p>离线演讲稿 · 内容来自当前 JSON / PPTX 备注。键盘 ← → 或空格翻页。</p>
<p id="timing-summary"></p><div class="toolbar">
<button id="prev" type="button" aria-label="上一页">← 上一页</button><select id="jump" aria-label="跳转页面"></select><button id="next" type="button" aria-label="下一页">下一页 →</button>
<button id="script-mode" type="button" aria-pressed="true">分页讲稿</button><button id="cards-mode" type="button" aria-pressed="false">提示卡</button><button id="all-mode" type="button" aria-pressed="false">全文</button>
<button id="start" type="button">开始计时</button><button id="reset" type="button">重置计时</button><span class="clock" aria-live="off">练习总计 <b id="total-clock">00:00</b> · 本页 <b id="page-clock">00:00</b></span>
</div></header><main><div id="missing" class="notice" hidden></div><div id="slides"></div></main>
<footer>文字估时与实际练习计时分别显示。未提供的稿件不补写；计时仅在当前浏览器页面运行。</footer>
<script id="deck-data" type="application/json">__DATA__</script>
<script>
"use strict";
const deck = JSON.parse(document.getElementById("deck-data").textContent);
const $ = id => document.getElementById(id);
const format = seconds => {const n=Math.max(0,Math.round(seconds)); return String(Math.floor(n/60)).padStart(2,"0")+":"+String(n%60).padStart(2,"0")};
let index=0, mode="script", running=false, last=0, total=0, page=0;
function elem(tag, cls, text){const node=document.createElement(tag);if(cls)node.className=cls;if(text!==undefined)node.textContent=text;return node}
const missing=[];
const cards=deck.slides.map((slide,i)=>{
 const node=elem("article","slide");node.dataset.slideId=slide.id;
 node.append(elem("h2","",(i+1)+". "+(slide.title||"未提供标题")));
 const t=slide.timing;node.append(elem("div","meta","ID "+slide.id+" · "+format(t.start_seconds)+"–"+format(t.end_seconds)+" · "+(t.basis==="storyboard_plan"?"故事板计划":"文字估计")+" "+format(t.scheduled_seconds)));
 node.append(elem("div",slide.notes?"notes":"notes empty",slide.notes||"缺少原始备注／口语稿，未自动补写。"));
 const cues=elem("ul","cues");
 if(slide.cues.length)slide.cues.forEach(c=>cues.append(elem("li","",c)));else cues.append(elem("li","empty","未提供提示卡。可切回分页讲稿阅读原备注。"));
 node.append(cues);
 if(slide.sources.length){const details=elem("details");details.append(elem("summary","","来源与依据"));const list=elem("ul");slide.sources.forEach(s=>list.append(elem("li","",s)));details.append(list);node.append(details)}
 if(!slide.notes)missing.push(i+1);
 $("slides").append(node);const option=elem("option","",(i+1)+". "+(slide.title||"未提供标题"));option.value=i;$("jump").append(option);return node;
});
if(missing.length){$("missing").hidden=false;$("missing").textContent="缺少原始备注／口语稿："+missing.join("、")+" 页。缺项已保留，未自动补写。"}
$("timing-summary").textContent="计划合计 "+format(deck.timing.scheduled_seconds)+"；文字估时 "+format(deck.timing.estimated_seconds)+"（中文 "+deck.timing.zh_characters_per_minute+" 字/分钟，英文 "+deck.timing.english_words_per_minute+" 词/分钟，单页停顿 "+deck.timing.pause_seconds_per_nonempty_slide+" 秒；不含未记录的互动与问答）。";
function tick(){const now=performance.now();if(running){const delta=now-last;total+=delta;page+=delta}last=now;$("total-clock").textContent=format(total/1000);$("page-clock").textContent=format(page/1000)}
function show(){document.body.dataset.mode=mode;cards.forEach((c,i)=>{c.hidden=mode!=="all"&&i!==index});$("jump").value=index;$("prev").disabled=index===0;$("next").disabled=index===cards.length-1;["script","cards","all"].forEach(m=>$(m+"-mode").setAttribute("aria-pressed",String(mode===m)))}
function navigate(to){const next=Math.max(0,Math.min(cards.length-1,to));if(next!==index){tick();index=next;page=0;$("page-clock").textContent="00:00"}show();if(mode==="all")cards[index].scrollIntoView({block:"start"})}
$("prev").onclick=()=>navigate(index-1);$("next").onclick=()=>navigate(index+1);$("jump").onchange=e=>navigate(Number(e.target.value));
["script","cards","all"].forEach(m=>$(m+"-mode").onclick=()=>{mode=m;show()});
$("start").onclick=()=>{tick();running=!running;last=performance.now();$("start").textContent=running?"暂停计时":(total>0?"继续计时":"开始计时")};
$("reset").onclick=()=>{running=false;total=0;page=0;last=performance.now();tick();$("start").textContent="开始计时"};
document.addEventListener("keydown",e=>{if(e.target.closest("button,input,select,textarea"))return;if(e.key==="ArrowRight"||e.key==="PageDown"||e.code==="Space"){e.preventDefault();navigate(index+1)}else if(e.key==="ArrowLeft"||e.key==="PageUp"){e.preventDefault();navigate(index-1)}});
show();setInterval(tick,250);
</script></body></html>
"""


def render_html(deck):
    payload = json.dumps(deck, ensure_ascii=False).replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e")
    values = {"__TITLE__": html.escape(deck["title"]), "__DATA__": payload}
    return re.sub(r"__TITLE__|__DATA__", lambda match: values[match.group(0)], HTML_TEMPLATE)


def render_docx(deck, path):
    from docx import Document
    from docx.shared import Pt
    document = Document()
    document.add_heading(deck["title"], 0)
    document.add_paragraph(f"计划总时长 {mmss(deck['timing']['scheduled_seconds'])}；文字估时 "
                           f"{mmss(deck['timing']['estimated_seconds'])}。估时不代表实际排练。")
    for pos, slide in enumerate(deck["slides"], 1):
        document.add_heading(f"{pos}. {slide['title'] or '未提供标题'}", 1)
        t = slide["timing"]
        document.add_paragraph(f"ID：{slide['id']} | {mmss(t['start_seconds'])}–{mmss(t['end_seconds'])}"
                               f" | {'故事板计划' if t['basis']=='storyboard_plan' else '文字估计'}")
        for paragraph in (slide["notes"] or "【缺少原始备注／口语稿；未自动补写】").split("\n"):
            document.add_paragraph(paragraph)
        if slide["cues"]:
            document.add_heading("提示卡", 2)
            for cue in slide["cues"]:
                document.add_paragraph(cue, style="List Bullet")
        if slide["sources"]:
            document.add_heading("来源", 2)
            for source in slide["sources"]:
                document.add_paragraph(source, style="List Bullet")
    document.styles["Normal"].font.size = Pt(12)
    document.save(str(path))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Current .json storyboard or .pptx deck")
    parser.add_argument("--out-dir", type=Path, help="Output directory, required unless --check-only")
    parser.add_argument("--pptx", type=Path, help="For JSON input: verify page order/id/title/notes against this current PPTX")
    parser.add_argument("--docx", action="store_true", help="Also export DOCX (requires python-docx)")
    parser.add_argument("--zh-cpm", type=float, default=230.0, help="Chinese characters/minute (default 230)")
    parser.add_argument("--en-wpm", type=float, default=130.0, help="English words/minute (default 130)")
    parser.add_argument("--pause-seconds", type=float, default=5.0, help="Extra pause seconds per non-empty page (default 5)")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacing existing generated files in the output directory")
    parser.add_argument("--check-only", action="store_true", help="Validate and report without creating files")
    args = parser.parse_args(argv)
    try:
        zh = numeric(args.zh_cpm, "--zh-cpm", positive=True)
        en = numeric(args.en_wpm, "--en-wpm", positive=True)
        pause = numeric(args.pause_seconds, "--pause-seconds")
        deck = read_input(args.input)
        sync = {"status": "not_requested"}
        if args.pptx:
            if args.input.suffix.lower() != ".json":
                raise ValueError("--pptx synchronization check is available only with JSON input")
            sync = check_sync(deck, read_pptx(args.pptx))
        findings = prepare(deck, zh, en, pause)
        report = {"status": "validated" if args.check_only else "exported", "input": str(args.input.resolve()),
                  "slide_count": len(deck["slides"]), "slide_ids": [s["id"] for s in deck["slides"]],
                  "sync": sync, "timing": deck["timing"], **findings,
                  "text_origin": "actual supplied notes; no LLM generation",
                  "outputs": []}
        if not args.check_only:
            if args.out_dir is None:
                raise ValueError("--out-dir is required unless --check-only")
            destinations = [args.out_dir / "speaker_notes.md", args.out_dir / "speaker_notes.html",
                            args.out_dir / "speaker_pack_report.json"]
            if args.docx:
                try:
                    import docx  # noqa: F401
                except ImportError as exc:
                    raise ValueError("--docx requires python-docx") from exc
                destinations.append(args.out_dir / "speaker_notes.docx")
            conflicts = [str(p) for p in destinations if p.exists()]
            if conflicts and not args.overwrite:
                raise ValueError("Existing outputs would be replaced. Use a new directory or explicit --overwrite:\n"
                                 + "\n".join(conflicts))
            if args.out_dir.exists() and not args.out_dir.is_dir():
                raise ValueError("--out-dir must be a directory")
            args.out_dir.mkdir(parents=True, exist_ok=True)
            report["outputs"] = [str(p.resolve()) for p in destinations]
            destinations[0].write_text(render_markdown(deck), encoding="utf-8")
            destinations[1].write_text(render_html(deck), encoding="utf-8")
            if args.docx:
                render_docx(deck, destinations[-1])
            destinations[2].write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except (ValueError, OSError, KeyError, json.JSONDecodeError) as exc:
        print(f"speaker_pack: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
