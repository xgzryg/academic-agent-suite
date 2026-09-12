"""Generate an editable PPTX from a presentation storyboard; never skip failed slides."""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

THEMES = {
    'academic': dict(bg='F7F9FC', panel='EAF0F7', ink='152D48', muted='506176', accent='147D92', line='CAD6E2'),
    'business': dict(bg='FFFFFF', panel='EEF3F8', ink='17304B', muted='52647B', accent='D68B2C', line='CFD9E3'),
    'teaching': dict(bg='FFFDF7', panel='F1F2E5', ink='283D32', muted='5D6A5B', accent='B46935', line='DADDC9'),
    'dark': dict(bg='142333', panel='20364B', ink='F5F8FC', muted='BCCDDC', accent='56D2C5', line='3C556B'),
    'minimal': dict(bg='FAFAFA', panel='EEEEEE', ink='242424', muted='626262', accent='7356AA', line='D1D1D1'),
}
TYPES = {'cover', 'section', 'agenda', 'content', 'two_column', 'table', 'chart', 'timeline', 'image', 'quote', 'summary', 'contact'}
CHARTS = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'horizontal_bar': XL_CHART_TYPE.BAR_CLUSTERED,
          'stacked_bar': XL_CHART_TYPE.COLUMN_STACKED, 'line': XL_CHART_TYPE.LINE_MARKERS, 'pie': XL_CHART_TYPE.PIE}


def strings(value):
    """Keep nested speaker-supplied items, including their order."""
    result = []
    for item in value:
        if isinstance(item, dict):
            result.append((str(item.get('t', '')), 0))
            result.extend((str(sub), 1) for sub in item.get('s', []))
        else:
            result.append((str(item), 0))
    return result


def validate(data, base: Path):
    if not isinstance(data, dict) or not isinstance(data.get('slides'), list) or not data['slides']:
        raise ValueError('The storyboard must contain a non-empty slides list.')
    if data.get('style', 'academic') not in THEMES:
        raise ValueError('Unknown style; choose: ' + ', '.join(THEMES))
    if data.get('aspect_ratio', '16:9') not in ('16:9', '4:3'):
        raise ValueError('Native layouts support 16:9 or 4:3. Use the SVG engine for custom canvases.')
    identifiers = set()
    for index, slide in enumerate(data['slides'], 1):
        if not isinstance(slide, dict):
            raise ValueError(f'Slide {index} must be an object.')
        ident = slide.get('id', f's{index:02d}')
        if not isinstance(ident, str) or not ident.strip() or ident in identifiers:
            raise ValueError(f'Slide {index}: id must be a unique non-empty string.')
        identifiers.add(ident)
        kind = slide.get('type', 'content')
        if kind not in TYPES:
            raise ValueError(f'Slide {ident}: unsupported type {kind!r}.')
        item_fields = {'content': ('items',), 'agenda': ('topics',), 'summary': ('points',),
                       'two_column': ('left', 'right'), 'timeline': ('milestones',)}
        for field in item_fields.get(kind, ()):
            if not isinstance(slide.get(field), list) or not slide[field]:
                raise ValueError(f'Slide {ident}: {field} must be a non-empty list.')
            for item in slide[field]:
                if not isinstance(item, (str, dict)):
                    raise ValueError(f'Slide {ident}: {field} items must be text or objects.')
                if isinstance(item, dict) and 's' in item and (not isinstance(item['s'], list) or
                        not all(isinstance(x, str) for x in item['s'])):
                    raise ValueError(f'Slide {ident}: nested s items must be a list of text.')
        for field in ('title', 'notes'):
            if field in slide and not isinstance(slide[field], str):
                raise ValueError(f'Slide {ident}: {field} must be text.')
        for field in ('sources', 'cues'):
            if field in slide and (not isinstance(slide[field], list) or not all(isinstance(x, str) for x in slide[field])):
                raise ValueError(f'Slide {ident}: {field} must be a list of text.')
        duration = slide.get('duration_seconds')
        if duration is not None and (isinstance(duration, bool) or not isinstance(duration, (int, float)) or not math.isfinite(duration) or duration < 0):
            raise ValueError(f'Slide {ident}: duration_seconds must be a finite non-negative number.')
        if kind == 'image':
            image_path = slide.get('image_path')
            if not image_path or not (base / image_path).is_file():
                raise FileNotFoundError(f'Slide {ident}: image_path is missing or unreadable: {image_path}')
            with Image.open(base / image_path) as img:
                img.verify()
            if slide.get('fit', 'contain') not in ('contain', 'cover'):
                raise ValueError(f'Slide {ident}: fit must be contain or cover.')
        if kind == 'table':
            headers, rows = slide.get('headers', []), slide.get('rows', [])
            if not isinstance(headers, list) or not headers or not isinstance(rows, list) or not rows or any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
                raise ValueError(f'Slide {ident}: table rows must match non-empty headers.')
        if kind == 'chart':
            categories, series = slide.get('categories', []), slide.get('series', [])
            chart_type = slide.get('chart_type', 'bar')
            if chart_type not in CHARTS or not isinstance(categories, list) or not categories or not isinstance(series, list) or not series:
                raise ValueError(f'Slide {ident}: provide a supported chart_type, categories and series.')
            for group in series:
                if not isinstance(group, dict) or not isinstance(group.get('values'), list):
                    raise ValueError(f'Slide {ident}: each series needs a values list.')
                values = group.get('values', [])
                if len(values) != len(categories):
                    raise ValueError(f'Slide {ident}: every series must match the category count.')
                if any(value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)) for value in values):
                    raise ValueError(f'Slide {ident}: chart values must be finite numbers or null for missing.')
            if chart_type == 'pie':
                if len(series) != 1 or any(v is None or v < 0 for v in series[0]['values']) or sum(series[0]['values']) <= 0:
                    raise ValueError(f'Slide {ident}: pie charts require one non-negative series with a positive total.')
    return data


class NativeDeck:
    def __init__(self, data, base):
        self.data, self.base = data, base
        self.prs = Presentation()
        self.w = 13.333333 if data.get('aspect_ratio', '16:9') == '16:9' else 10.0
        self.prs.slide_width, self.prs.slide_height = Inches(self.w), Inches(7.5)
        self.theme = dict(THEMES[data.get('style', 'academic')])
        self.font = data.get('font', 'Aptos')
        self.cjk_font = data.get('cjk_font', 'Microsoft YaHei')
        self.warnings = []

    def x(self, value):
        return Inches(value * self.w / 13.333333)

    def color(self, key):
        return RGBColor.from_string(self.theme.get(key, key))

    def shape(self, slide, x, y, w, h, color, rounded=False):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                       self.x(x), Inches(y), self.x(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.color(color)
        shape.line.fill.background()
        return shape

    def apply_font(self, paragraph, size, color, bold=False):
        paragraph.font.name = self.font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = self.color(color)
        props = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
        ea = props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
        if ea is None:
            ea = OxmlElement('a:ea')
            props.append(ea)
        ea.set('typeface', self.cjk_font)

    def text(self, slide, x, y, w, h, content, size=24, color='ink', bold=False, align=PP_ALIGN.LEFT, name=None):
        box = slide.shapes.add_textbox(self.x(x), Inches(y), self.x(w), Inches(h))
        if name:
            box.name = name
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        rows = content if isinstance(content, list) else str(content).split('\n')
        for index, row in enumerate(rows):
            text, level = row if isinstance(row, tuple) else (str(row), 0)
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.text = ('  ' if level else '') + text
            para.alignment = align
            para.space_after = Pt(12 if len(rows) < 7 else 7)
            self.apply_font(para, max(8, size - level * 3), color, bold)
        # This is a preparation warning, not a claim of actual PowerPoint overflow detection.
        if sum(len(str(row)) for row in rows) > w * h * 21:
            self.warnings.append(f'Dense text in {name or "a text block"}; inspect the rendered slide and split content if needed.')
        return box

    def header(self, slide, item, ident):
        self.shape(slide, 0.0, 0.0, 13.333, 0.08, 'accent')
        self.text(slide, 0.68, 0.39, 11.95, 0.80, item.get('title', ''), 29, bold=True, name='ppa-title:' + ident)
        self.shape(slide, 0.68, 1.29, 11.97, 0.015, 'line')

    def finish(self, slide, item, ident, number):
        if item.get('sources'):
            self.text(slide, 0.68, 6.52, 11.65, 0.37, ' | '.join(item['sources']), 12, 'muted')
        footer = self.data.get('footer', '')
        if footer:
            self.text(slide, 0.68, 7.02, 10.7, 0.23, footer, 10, 'muted')
        if self.data.get('page_numbers', True):
            self.text(slide, 11.9, 6.99, 0.72, 0.25, f'{number:02d}', 11, 'muted', align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = item.get('notes', '')

    def add(self, item, index):
        ident = item.get('id', f's{index:02d}')
        kind = item.get('type', 'content')
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.color('bg')
        if kind in ('cover', 'section'):
            self.shape(slide, 0.65, 1.36, 0.11, 4.53, 'accent')
            if item.get('kicker'):
                self.text(slide, 1.02, 1.20, 10.8, 0.45, item['kicker'], 15, 'muted')
            self.text(slide, 1.02, 2.06, 11.0, 2.0, item.get('title', ''), 40 if kind == 'cover' else 36,
                      bold=True, name='ppa-title:' + ident)
            if item.get('subtitle'):
                self.text(slide, 1.05, 4.41, 10.9, 1.0, item['subtitle'], 23, 'muted')
            meta = ' · '.join(str(item[key]) for key in ('author', 'date') if item.get(key))
            if meta:
                self.text(slide, 1.05, 5.9, 10.9, 0.42, meta, 16, 'muted')
        else:
            self.header(slide, item, ident)
            if kind in ('content', 'agenda', 'summary'):
                key = {'content': 'items', 'agenda': 'topics', 'summary': 'points'}[kind]
                rows = strings(item.get(key, []))
                if kind == 'agenda':
                    rows = [(f'{i:02d}  {text}', level) for i, (text, level) in enumerate(rows, 1)]
                self.text(slide, 0.92, 1.75, 11.45, 3.85 if item.get('conclusion') else 4.45,
                          rows, 24 if len(rows) <= 6 else 20)
                if item.get('conclusion'):
                    self.shape(slide, 0.75, 5.75, 11.84, 0.58, 'panel', rounded=True)
                    self.text(slide, 1.0, 5.88, 11.35, 0.40, item['conclusion'], 19, 'ink', True)
            elif kind == 'two_column':
                for n, key in enumerate(('left', 'right')):
                    x = 0.72 + n * 6.09
                    self.shape(slide, x, 1.63, 5.80, 4.62, 'panel', rounded=True)
                    label = item.get(key + '_title', '')
                    top = 2.39 if label else 1.98
                    if label:
                        self.text(slide, x + 0.28, 1.91, 5.20, 0.45, label, 21, 'accent', True)
                    self.text(slide, x + 0.28, top, 5.20, 3.60, strings(item.get(key, [])), 21)
            elif kind == 'table':
                headers, rows = item['headers'], item['rows']
                table = slide.shapes.add_table(len(rows) + 1, len(headers), self.x(0.75), Inches(1.65),
                                               self.x(11.84), Inches(4.62)).table
                for ri, row in enumerate([headers, *rows]):
                    for ci, value in enumerate(row):
                        cell = table.cell(ri, ci)
                        cell.text = '—' if value is None else str(value)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = self.color('ink' if ri == 0 else ('panel' if ri % 2 else 'bg'))
                        for para in cell.text_frame.paragraphs:
                            self.apply_font(para, 18 if len(rows) <= 8 else 15, 'bg' if ri == 0 else 'ink', ri == 0)
                if any(value is None for row in rows for value in row):
                    self.text(slide, 0.76, 6.30, 8.0, 0.21, '— = missing / 缺失', 11, 'muted')
                if len(rows) > 8 or len(headers) > 6:
                    self.warnings.append(f'{ident}: dense table; split into more slides for projection.')
            elif kind == 'chart':
                cd = CategoryChartData()
                cd.categories = item['categories']
                for series in item['series']:
                    cd.add_series(str(series.get('name', '')), series['values'])
                chart = slide.shapes.add_chart(CHARTS[item.get('chart_type', 'bar')], self.x(0.80), Inches(1.62),
                                                self.x(11.72), Inches(4.64), cd).chart
                chart.has_legend = len(item['series']) > 1 or item.get('chart_type') == 'pie'
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                    chart.legend.font.size = Pt(16)
                chart.font.name = self.font
                chart.font.size = Pt(16)
                chart.font.color.rgb = self.color('ink')
                if chart.has_legend:
                    chart.legend.font.color.rgb = self.color('ink')
                for n, series in enumerate(chart.series):
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = self.color(('accent', 'muted', 'ink')[n % 3])
                    if item.get('chart_type') == 'line':
                        series.format.line.color.rgb = self.color(('accent', 'muted', 'ink')[n % 3])
                if item.get('chart_type') == 'pie':
                    palette = ('147D92', 'D59B39', '526A99', 'A86778', '669968', '8776A4')
                    for n, point in enumerate(chart.series[0].points):
                        point.format.fill.solid()
                        point.format.fill.fore_color.rgb = RGBColor.from_string(palette[n % len(palette)])
                        point.format.line.color.rgb = self.color('bg')
                if item.get('chart_type') != 'pie':
                    chart.category_axis.tick_labels.font.size = Pt(15)
                    chart.value_axis.tick_labels.font.size = Pt(14)
                    chart.category_axis.tick_labels.font.color.rgb = self.color('ink')
                    chart.value_axis.tick_labels.font.color.rgb = self.color('ink')
                    for axis in (chart.category_axis, chart.value_axis):
                        axis.format.line.color.rgb = self.color('muted')
                    if chart.value_axis.has_major_gridlines:
                        chart.value_axis.major_gridlines.format.line.color.rgb = self.color('line')
                    if item.get('y_label'):
                        chart.value_axis.has_title = True
                        chart.value_axis.axis_title.text_frame.text = item['y_label']
                        for para in chart.value_axis.axis_title.text_frame.paragraphs:
                            self.apply_font(para, 15, 'ink')
            elif kind == 'timeline':
                points = item.get('milestones', [])
                if not points:
                    raise ValueError(f'Slide {ident}: milestones are required.')
                self.shape(slide, 1.0, 3.25, 11.3, 0.035, 'line')
                width = 11.3 / len(points)
                for n, point in enumerate(points):
                    point = point if isinstance(point, dict) else {'label': str(point)}
                    cx = 1.0 + width * (n + 0.5)
                    self.shape(slide, cx - 0.10, 3.14, 0.20, 0.24, 'accent', rounded=True)
                    self.text(slide, cx - width / 2 + 0.1, 2.13, width - 0.2, 0.9,
                              point.get('label', ''), 21, 'ink', True, PP_ALIGN.CENTER)
                    self.text(slide, cx - width / 2 + 0.1, 3.67, width - 0.2, 1.85,
                              point.get('desc', ''), 18, 'muted', align=PP_ALIGN.CENTER)
            elif kind == 'image':
                path = (self.base / item['image_path']).resolve()
                bx, by, bw, bh = self.x(0.85), Inches(1.57), self.x(11.63), Inches(4.53)
                if item.get('fit', 'contain') == 'cover':
                    picture = slide.shapes.add_picture(str(path), bx, by)
                    scale = max(bw / picture.width, bh / picture.height)
                    pw, ph = picture.width * scale, picture.height * scale
                    picture.crop_left = picture.crop_right = max(0, (pw - bw) / (2 * pw))
                    picture.crop_top = picture.crop_bottom = max(0, (ph - bh) / (2 * ph))
                    picture.width, picture.height = bw, bh
                else:
                    with Image.open(path) as img:
                        iw, ih = img.size
                    factor = min(bw / iw, bh / ih)
                    pw, ph = int(iw * factor), int(ih * factor)
                    slide.shapes.add_picture(str(path), bx + (bw - pw) // 2, by + (bh - ph) // 2, width=pw, height=ph)
                if item.get('caption'):
                    self.text(slide, 0.85, 6.15, 11.63, 0.32, item['caption'], 14, 'muted', align=PP_ALIGN.CENTER)
            elif kind == 'quote':
                self.shape(slide, 0.85, 1.75, 11.63, 4.31, 'panel', rounded=True)
                self.shape(slide, 0.85, 1.75, 0.09, 4.31, 'accent')
                self.text(slide, 1.30, 2.29, 10.63, 2.55, item.get('quote', ''), 29)
                if item.get('attribution'):
                    self.text(slide, 1.30, 5.16, 10.63, 0.52, item['attribution'], 18, 'muted', align=PP_ALIGN.RIGHT)
            elif kind == 'contact':
                self.text(slide, 1.0, 2.20, 11.3, 2.60,
                          item.get('info', ''), 27, align=PP_ALIGN.CENTER)
        self.finish(slide, item, ident, index)

    def build(self):
        for index, item in enumerate(self.data['slides'], 1):
            self.add(item, index)
        return self.prs


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('storyboard', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--overwrite', action='store_true', help='Use only when replacement of this output is authorized.')
    args = parser.parse_args()
    if args.output.suffix.lower() != '.pptx':
        parser.error('--output must end in .pptx')
    if args.output.exists() and not args.overwrite:
        raise FileExistsError('Output already exists. Choose a new path or use an authorized --overwrite.')
    data = json.loads(args.storyboard.read_text(encoding='utf-8-sig'))
    validate(data, args.storyboard.resolve().parent)
    generator = NativeDeck(data, args.storyboard.resolve().parent)
    presentation = generator.build()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)
    print(json.dumps({'output': str(args.output.resolve()), 'slides': len(presentation.slides),
                      'notes_present': sum(bool(s.get('notes', '').strip()) for s in data['slides']),
                      'layout_warnings': generator.warnings, 'visual_review': 'required'}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except (ValueError, OSError, TypeError, KeyError) as exc:
        print(f'Presentation not generated: {exc}', file=sys.stderr)
        raise SystemExit(2)
