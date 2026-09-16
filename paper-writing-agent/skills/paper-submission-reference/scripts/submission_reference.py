#!/usr/bin/env python3
"""Create, validate, and render the first-submission information master workbook."""

from __future__ import annotations

import argparse
from datetime import date
import json
from pathlib import Path
import sys
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


VERSION = "1.0"
STATUS_VALUES = ["已确认", "抽取待确认", "缺失", "不适用"]
VISIBILITY_VALUES = ["论文公开", "仅投稿系统", "内部使用"]
YES_NO_VALUES = ["是", "否", "不适用", "尚不确定"]
READY_VALUES = ["已准备", "待准备", "不适用", "尚不确定"]

FIELD_VALUE_OPTIONS: dict[str, list[str]] = {
    "funding_status": YES_NO_VALUES,
    "conflict_status": YES_NO_VALUES,
    "ai_use_status": YES_NO_VALUES,
    "originality_confirmed": YES_NO_VALUES,
    "simultaneous_submission_confirmed": YES_NO_VALUES,
    "preprint_status": YES_NO_VALUES,
    "related_manuscript_status": YES_NO_VALUES,
    "third_party_assistance_status": YES_NO_VALUES,
    "third_party_materials_applicable": YES_NO_VALUES,
    "human_research_applicable": YES_NO_VALUES,
    "animal_research_applicable": YES_NO_VALUES,
    "trial_registration_applicable": YES_NO_VALUES,
    "code_generated_applicable": YES_NO_VALUES,
    "data_access_type": ["公开仓库", "受控仓库", "合理申请", "受限申请", "因隐私不公开", "不适用", "尚不确定"],
    "ethics_approval_status": ["已批准", "已豁免", "不适用", "尚不确定"],
    "informed_consent_status": ["已取得", "已豁免", "不适用", "尚不确定"],
    "publication_consent_status": ["已取得", "不适用", "尚不确定"],
    "animal_ethics_status": ["已批准", "已豁免", "不适用", "尚不确定"],
    "third_party_permissions_status": ["已取得", "申请中", "不适用", "尚不确定"],
}

SHEET_STATUS = "使用说明与状态"
SHEET_MANUSCRIPT = "稿件信息"
SHEET_AUTHORS = "作者信息"
SHEET_AFFILIATIONS = "作者单位"
SHEET_CONTRIBUTIONS = "贡献与确认"
SHEET_DECLARATIONS = "声明"
SHEET_ETHICS = "伦理与注册"
SHEET_DATA = "数据代码材料"
SHEET_REVIEWERS = "编辑与审稿人"
SHEET_FILES = "上传文件清单"
SHEET_ISSUES = "缺项与待确认"
SHEET_METADATA = "字段元数据"

FIELD_HEADERS = [
    "字段ID",
    "中文字段",
    "English field",
    "值",
    "状态",
    "来源",
    "可见性",
    "备注",
    "最后确认日期",
    "关键项",
    "适用性",
]

AUTHOR_HEADERS = [
    "作者ID", "作者顺序", "名/Given name", "中间名/Middle name", "姓/Family name",
    "投稿显示名", "学位", "职位", "称谓", "邮箱", "ORCID", "性别或系统选项",
    "国家/地区", "单位ID", "通讯作者", "共同第一作者", "共同资深作者", "责任作者/保证人",
    "提交人", "电话", "记录状态", "来源", "可见性", "备注", "最后确认日期",
]

AFFILIATION_HEADERS = [
    "单位ID", "机构", "科室/部门", "城市", "省/州", "邮编", "国家/地区", "关联作者ID",
    "记录状态", "来源", "可见性", "备注", "最后确认日期",
]

CREDIT_ROLES = [
    "Conceptualization", "Data curation", "Formal analysis", "Funding acquisition",
    "Investigation", "Methodology", "Project administration", "Resources", "Software",
    "Supervision", "Validation", "Visualization", "Writing – original draft",
    "Writing – review & editing",
]

CONTRIBUTION_HEADERS = [
    "作者ID", *CREDIT_ROLES, "作者资格", "批准最终稿", "同意投稿", "责任确认",
    "记录状态", "来源", "可见性", "备注", "最后确认日期",
]

REVIEWER_HEADERS = [
    "记录ID", "类型", "姓名", "机构", "邮箱", "理由", "关系披露", "记录状态",
    "来源", "可见性", "备注", "最后确认日期",
]

FILE_HEADERS = [
    "文件ID", "文件类型", "文件名", "版本", "是否匿名", "是否必需", "准备状态",
    "记录状态", "来源", "可见性", "备注", "最后确认日期",
]

ISSUE_HEADERS = [
    "编号", "工作表", "对象/字段", "当前值", "状态", "是否阻塞最终版", "原因", "下一步",
]

METADATA_HEADERS = [
    "工作表", "对象ID", "字段名", "当前值", "状态", "来源", "可见性", "备注", "最后确认日期",
]


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def field(
    field_id: str,
    zh: str,
    en: str,
    *,
    critical: bool = False,
    visibility: str = "论文公开",
    applicability: str = "全部",
    allow_na: bool = False,
    default_value: str = "",
    default_status: str = "缺失",
    default_source: str = "",
) -> dict[str, Any]:
    return {
        "id": field_id,
        "zh": zh,
        "en": en,
        "critical": critical,
        "visibility": visibility,
        "applicability": applicability,
        "allow_na": allow_na,
        "default_value": default_value,
        "default_status": default_status,
        "default_source": default_source,
    }


FIELD_SPECS: dict[str, list[dict[str, Any]]] = {
    SHEET_MANUSCRIPT: [
        field("manuscript_title", "稿件完整标题", "Full manuscript title", critical=True),
        field("short_title", "短标题", "Short title"),
        field("article_type", "文章类型", "Article type", critical=True),
        field("abstract", "摘要", "Abstract"),
        field("keywords", "关键词", "Keywords"),
        field("word_count", "正文字数", "Main-text word count", visibility="内部使用"),
        field("table_count", "表格数量", "Number of tables", visibility="内部使用"),
        field("figure_count", "图形数量", "Number of figures", visibility="内部使用"),
        field("supplementary_count", "补充文件数量", "Number of supplementary files", visibility="内部使用"),
        field("target_journal", "目标期刊（如已确定）", "Target journal"),
        field("manuscript_id", "稿件编号（如已有）", "Manuscript ID", visibility="仅投稿系统"),
        field(
            "submission_stage", "投稿阶段", "Submission stage", critical=True,
            visibility="内部使用", default_value="首次投稿", default_status="已确认",
            default_source="本功能固定范围",
        ),
        field("submitter_author_id", "提交人作者ID", "Submitting author ID", critical=True, visibility="仅投稿系统"),
        field("corresponding_author_ids", "通讯作者ID", "Corresponding author ID(s)", critical=True, visibility="仅投稿系统"),
        field("manuscript_language", "稿件语言", "Manuscript language"),
    ],
    SHEET_DECLARATIONS: [
        field("funding_status", "是否获得基金或其他资助", "Funding received", critical=True, visibility="论文公开"),
        field("funding_statement_en", "基金英文声明", "Funding statement", critical=True, applicability="funding_status=是", allow_na=True),
        field("funder_role_en", "资助方作用英文声明", "Role of the funder", critical=True, applicability="funding_status=是", allow_na=True),
        field("conflict_status", "是否存在利益冲突", "Competing interests present", critical=True),
        field("conflict_statement_en", "利益冲突英文声明", "Competing interests statement", critical=True, allow_na=False),
        field("acknowledgements_en", "致谢英文正文", "Acknowledgements"),
        field("ai_use_status", "是否使用生成式AI或自动化辅助", "Use of generative AI or automated assistance", critical=True, visibility="仅投稿系统"),
        field("ai_use_statement_en", "AI使用英文声明", "AI use statement", critical=True, applicability="ai_use_status=是", allow_na=True),
        field("originality_confirmed", "原创性已由作者确认", "Originality confirmed", critical=True, visibility="仅投稿系统"),
        field("simultaneous_submission_confirmed", "非一稿多投已由作者确认", "No simultaneous submission confirmed", critical=True, visibility="仅投稿系统"),
        field("preprint_status", "是否存在预印本", "Preprint available", critical=True),
        field("preprint_statement_en", "预印本英文说明", "Preprint statement", critical=True, applicability="preprint_status=是", allow_na=True),
        field("preprint_url", "预印本链接或DOI", "Preprint URL or DOI", critical=True, applicability="preprint_status=是", allow_na=True),
        field("related_manuscript_status", "是否有相关稿件", "Related manuscript exists", critical=True, visibility="仅投稿系统"),
        field("related_manuscript_details_en", "相关稿件英文说明", "Related manuscript details", critical=True, applicability="related_manuscript_status=是", visibility="仅投稿系统", allow_na=True),
        field("third_party_assistance_status", "是否有第三方写作或技术协助", "Third-party assistance", visibility="仅投稿系统"),
        field("third_party_assistance_details_en", "第三方协助英文说明", "Third-party assistance details", applicability="third_party_assistance_status=是", visibility="仅投稿系统", allow_na=True),
        field("third_party_materials_applicable", "是否使用需许可的第三方材料", "Third-party permissions applicable", critical=True, visibility="仅投稿系统"),
        field("third_party_permissions_status", "第三方许可状态", "Third-party permission status", critical=True, applicability="third_party_materials_applicable=是", visibility="仅投稿系统", allow_na=True),
    ],
    SHEET_ETHICS: [
        field("human_research_applicable", "是否涉及人体参与者、样本或可识别数据", "Human research applicable", critical=True),
        field("ethics_approval_status", "人体伦理批准或豁免状态", "Ethics approval or exemption status", critical=True, applicability="human_research_applicable=是", allow_na=True),
        field("ethics_committee", "伦理委员会全名", "Ethics committee", critical=True, applicability="human_research_applicable=是", allow_na=True),
        field("ethics_approval_number", "人体伦理批准号", "Ethics approval number", critical=True, applicability="human_research_applicable=是", allow_na=True),
        field("informed_consent_status", "知情同意状态", "Informed consent status", critical=True, applicability="human_research_applicable=是", allow_na=True),
        field("publication_consent_status", "发表同意状态", "Consent for publication status", critical=True, applicability="human_research_applicable=是", allow_na=True),
        field("ethics_statement_en", "人体伦理英文声明", "Ethics statement", critical=True, applicability="human_research_applicable=是", allow_na=True),
        field("animal_research_applicable", "是否涉及动物研究", "Animal research applicable", critical=True),
        field("animal_ethics_status", "动物伦理批准或豁免状态", "Animal ethics approval status", critical=True, applicability="animal_research_applicable=是", allow_na=True),
        field("animal_committee", "动物伦理机构全名", "Animal ethics committee", critical=True, applicability="animal_research_applicable=是", allow_na=True),
        field("animal_approval_number", "动物伦理批准号", "Animal ethics approval number", critical=True, applicability="animal_research_applicable=是", allow_na=True),
        field("animal_ethics_statement_en", "动物伦理英文声明", "Animal ethics statement", critical=True, applicability="animal_research_applicable=是", allow_na=True),
        field("trial_registration_applicable", "是否需要试验或研究注册", "Registration applicable", critical=True),
        field("registry_name", "注册平台", "Registry", critical=True, applicability="trial_registration_applicable=是", allow_na=True),
        field("registration_number", "注册号", "Registration number", critical=True, applicability="trial_registration_applicable=是", allow_na=True),
        field("registration_date", "注册日期", "Registration date", applicability="trial_registration_applicable=是"),
        field("reporting_guideline", "采用的报告规范", "Reporting guideline"),
    ],
    SHEET_DATA: [
        field("data_access_type", "数据访问方式", "Data access type", critical=True),
        field("data_availability_statement_en", "数据可用性英文声明", "Data availability statement", critical=True),
        field("data_repository", "数据仓库", "Data repository", applicability="data_access_type=公开仓库|受控仓库"),
        field("data_url_or_id", "数据链接或持久标识符", "Data URL or persistent identifier", applicability="data_access_type=公开仓库|受控仓库"),
        field("data_restriction_reason_en", "数据限制英文说明", "Reason for data restriction", applicability="data_access_type=受限申请|因隐私不公开"),
        field("data_contact", "数据访问联系人", "Data access contact", visibility="仅投稿系统"),
        field("code_generated_applicable", "本研究是否产生分析代码或软件", "Code or software generated", critical=True),
        field("code_availability_statement_en", "代码可用性英文声明", "Code availability statement", critical=True, applicability="code_generated_applicable=是", allow_na=True),
        field("code_repository", "代码仓库", "Code repository", applicability="code_generated_applicable=是"),
        field("code_url_or_version", "代码链接或版本", "Code URL or version", applicability="code_generated_applicable=是"),
        field("materials_availability_statement_en", "材料可用性英文声明", "Materials availability statement"),
        field("materials_contact", "材料联系人", "Materials contact", visibility="仅投稿系统"),
    ],
}


def today_text() -> str:
    return date.today().isoformat()


def as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_yes(value: Any) -> str:
    text = as_text(value).lower()
    if text in {"yes", "y", "true", "1", "是"}:
        return "是"
    if text in {"no", "n", "false", "0", "否"}:
        return "否"
    return as_text(value)


def style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def style_sheet(ws, widths: dict[str, float] | None = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for col, width in widths.items():
            ws.column_dimensions[col].width = width


def add_list_validation(ws, column: int, values: list[str], start: int = 2, end: int = 2000) -> None:
    quoted = '"' + ",".join(values) + '"'
    validation = DataValidation(type="list", formula1=quoted, allow_blank=True)
    validation.error = "请从下拉列表中选择。"
    validation.errorTitle = "值不在允许范围"
    ws.add_data_validation(validation)
    letter = get_column_letter(column)
    validation.add(f"{letter}{start}:{letter}{end}")


def append_headers(ws, headers: list[str]) -> None:
    ws.append(headers)
    style_header(ws)


def create_status_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet(SHEET_STATUS)
    ws.append(["项目", "当前值", "说明"])
    rows = [
        ["模板版本", VERSION, "投稿信息总表结构版本"],
        ["生成日期", today_text(), "最近一次由脚本生成或验证的日期"],
        ["文档状态", "草稿版", "零阻塞项时自动更新为最终版"],
        ["阻塞项数量", 0, "关键字段未确认、不允许不适用或结构缺失"],
        ["非阻塞待办数量", 0, "可选字段、推荐审稿人或文件准备事项"],
        ["隐私提示", "含个人和投稿系统信息", "不要记录密码、验证码、令牌、银行卡或支付凭据"],
        ["资料源", "Excel 主表", "Word 仅由本工作簿生成，不作为反向更新来源"],
        ["适用范围", "跨期刊首次投稿通用资料", "不替代目标期刊现行指南"],
    ]
    for row in rows:
        ws.append(row)
    style_header(ws)
    style_sheet(ws, {"A": 22, "B": 32, "C": 72})
    ws.auto_filter.ref = f"A1:C{ws.max_row}"


def create_field_sheet(wb: Workbook, name: str, specs: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    append_headers(ws, FIELD_HEADERS)
    for spec in specs:
        ws.append([
            spec["id"], spec["zh"], spec["en"], spec["default_value"], spec["default_status"],
            spec["default_source"], spec["visibility"], "", "", "是" if spec["critical"] else "否",
            spec["applicability"],
        ])
        if spec["id"] in FIELD_VALUE_OPTIONS:
            add_list_validation(ws, 4, FIELD_VALUE_OPTIONS[spec["id"]], start=ws.max_row, end=ws.max_row)
    add_list_validation(ws, 5, STATUS_VALUES)
    add_list_validation(ws, 7, VISIBILITY_VALUES)
    style_sheet(ws, {
        "A": 30, "B": 28, "C": 31, "D": 58, "E": 15, "F": 24, "G": 16,
        "H": 42, "I": 16, "J": 12, "K": 30,
    })


def create_wide_sheet(
    wb: Workbook,
    name: str,
    headers: list[str],
    widths: dict[str, float],
    *,
    status_header: str = "记录状态",
    visibility_header: str = "可见性",
) -> None:
    ws = wb.create_sheet(name)
    append_headers(ws, headers)
    if status_header in headers:
        add_list_validation(ws, headers.index(status_header) + 1, STATUS_VALUES)
    if visibility_header in headers:
        add_list_validation(ws, headers.index(visibility_header) + 1, VISIBILITY_VALUES)
    style_sheet(ws, widths)


def create_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    create_status_sheet(wb)
    create_field_sheet(wb, SHEET_MANUSCRIPT, FIELD_SPECS[SHEET_MANUSCRIPT])

    create_wide_sheet(wb, SHEET_AUTHORS, AUTHOR_HEADERS, {
        "A": 12, "B": 10, "C": 18, "D": 18, "E": 18, "F": 22, "G": 14, "H": 18,
        "I": 14, "J": 28, "K": 24, "L": 20, "M": 18, "N": 18, "O": 14, "P": 15,
        "Q": 15, "R": 18, "S": 12, "T": 20, "U": 15, "V": 22, "W": 16, "X": 38, "Y": 16,
    })
    for header in ["通讯作者", "共同第一作者", "共同资深作者", "责任作者/保证人", "提交人"]:
        add_list_validation(wb[SHEET_AUTHORS], AUTHOR_HEADERS.index(header) + 1, YES_NO_VALUES)

    create_wide_sheet(wb, SHEET_AFFILIATIONS, AFFILIATION_HEADERS, {
        "A": 12, "B": 34, "C": 30, "D": 18, "E": 18, "F": 14, "G": 18, "H": 24,
        "I": 15, "J": 22, "K": 16, "L": 38, "M": 16,
    })
    create_wide_sheet(wb, SHEET_CONTRIBUTIONS, CONTRIBUTION_HEADERS, {
        "A": 12, **{get_column_letter(i): 18 for i in range(2, 16)},
        "P": 14, "Q": 14, "R": 14, "S": 14, "T": 15, "U": 22, "V": 16, "W": 38, "X": 16,
    })
    for header in CREDIT_ROLES + ["作者资格", "批准最终稿", "同意投稿", "责任确认"]:
        add_list_validation(wb[SHEET_CONTRIBUTIONS], CONTRIBUTION_HEADERS.index(header) + 1, YES_NO_VALUES)

    for sheet in [SHEET_DECLARATIONS, SHEET_ETHICS, SHEET_DATA]:
        create_field_sheet(wb, sheet, FIELD_SPECS[sheet])

    create_wide_sheet(wb, SHEET_REVIEWERS, REVIEWER_HEADERS, {
        "A": 12, "B": 18, "C": 22, "D": 34, "E": 28, "F": 38, "G": 34,
        "H": 15, "I": 22, "J": 16, "K": 38, "L": 16,
    })
    create_wide_sheet(wb, SHEET_FILES, FILE_HEADERS, {
        "A": 12, "B": 22, "C": 42, "D": 14, "E": 14, "F": 14, "G": 16,
        "H": 15, "I": 22, "J": 16, "K": 38, "L": 16,
    })
    for header in ["是否匿名", "是否必需"]:
        add_list_validation(wb[SHEET_FILES], FILE_HEADERS.index(header) + 1, YES_NO_VALUES)
    add_list_validation(wb[SHEET_FILES], FILE_HEADERS.index("准备状态") + 1, READY_VALUES)

    ws = wb.create_sheet(SHEET_ISSUES)
    append_headers(ws, ISSUE_HEADERS)
    style_sheet(ws, {"A": 8, "B": 20, "C": 34, "D": 42, "E": 16, "F": 18, "G": 52, "H": 46})

    metadata = wb.create_sheet(SHEET_METADATA)
    append_headers(metadata, METADATA_HEADERS)
    style_sheet(metadata, {"A": 20, "B": 18, "C": 30, "D": 45, "E": 16, "F": 24, "G": 16, "H": 42, "I": 16})
    metadata.sheet_state = "hidden"
    wb.active = 0
    return wb


def header_map(ws) -> dict[str, int]:
    return {as_text(cell.value): index for index, cell in enumerate(ws[1], start=1) if as_text(cell.value)}


def append_mapping_row(ws, mapping: dict[str, Any]) -> list[str]:
    headers = header_map(ws)
    unknown = sorted(set(mapping) - set(headers))
    row = [""] * len(headers)
    for key, value in mapping.items():
        if key in headers:
            row[headers[key] - 1] = value
    ws.append(row)
    return unknown


def apply_seed(wb: Workbook, seed: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    field_rows: dict[str, tuple[Any, int, dict[str, int]]] = {}
    for sheet in FIELD_SPECS:
        ws = wb[sheet]
        headers = header_map(ws)
        for row in range(2, ws.max_row + 1):
            field_id = as_text(ws.cell(row, headers["字段ID"]).value)
            field_rows[field_id] = (ws, row, headers)

    for field_id, supplied in (seed.get("fields") or {}).items():
        if field_id not in field_rows:
            warnings.append(f"Unknown field ID ignored: {field_id}")
            continue
        if not isinstance(supplied, dict):
            supplied = {"value": supplied}
        ws, row, headers = field_rows[field_id]
        key_map = {
            "value": "值", "status": "状态", "source": "来源", "visibility": "可见性",
            "note": "备注", "confirmed_date": "最后确认日期",
        }
        for source_key, header in key_map.items():
            if source_key in supplied:
                ws.cell(row, headers[header]).value = supplied[source_key]
        if "status" not in supplied:
            ws.cell(row, headers["状态"]).value = "抽取待确认"
        if "source" not in supplied:
            ws.cell(row, headers["来源"]).value = "材料抽取"

    entity_map = {
        "authors": SHEET_AUTHORS,
        "affiliations": SHEET_AFFILIATIONS,
        "contributions": SHEET_CONTRIBUTIONS,
        "reviewers": SHEET_REVIEWERS,
        "files": SHEET_FILES,
    }
    for key, sheet in entity_map.items():
        records = seed.get(key) or []
        if not isinstance(records, list):
            warnings.append(f"Seed key {key} must be a list; ignored.")
            continue
        for position, record in enumerate(records, start=1):
            if not isinstance(record, dict):
                warnings.append(f"Seed {key}[{position}] is not an object; ignored.")
                continue
            unknown = append_mapping_row(wb[sheet], record)
            for name in unknown:
                warnings.append(f"Unknown {sheet} column ignored: {name}")
    return warnings


def get_field_records(wb) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    records: dict[str, dict[str, Any]] = {}
    specs_by_id: dict[str, dict[str, Any]] = {}
    for sheet, specs in FIELD_SPECS.items():
        ws = wb[sheet]
        headers = header_map(ws)
        for spec in specs:
            specs_by_id[spec["id"]] = {**spec, "sheet": sheet}
        for row in range(2, ws.max_row + 1):
            field_id = as_text(ws.cell(row, headers["字段ID"]).value)
            if not field_id:
                continue
            records[field_id] = {
                "sheet": sheet,
                "field_id": field_id,
                "zh": as_text(ws.cell(row, headers["中文字段"]).value),
                "en": as_text(ws.cell(row, headers["English field"]).value),
                "value": as_text(ws.cell(row, headers["值"]).value),
                "status": as_text(ws.cell(row, headers["状态"]).value),
                "source": as_text(ws.cell(row, headers["来源"]).value),
                "visibility": as_text(ws.cell(row, headers["可见性"]).value),
                "note": as_text(ws.cell(row, headers["备注"]).value),
                "confirmed_date": as_text(ws.cell(row, headers["最后确认日期"]).value),
                "critical": as_text(ws.cell(row, headers["关键项"]).value) == "是",
                "applicability": as_text(ws.cell(row, headers["适用性"]).value) or "全部",
            }
    return records, specs_by_id


def condition_active(condition: str, values: dict[str, str]) -> bool:
    if not condition or condition == "全部":
        return True
    if "=" not in condition:
        return True
    field_id, expected = condition.split("=", 1)
    allowed = {part.strip() for part in expected.split("|")}
    return values.get(field_id.strip(), "") in allowed


def rows_as_dicts(ws) -> list[dict[str, Any]]:
    headers = [as_text(cell.value) for cell in ws[1]]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(as_text(value) for value in values):
            continue
        rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers)})
    return rows


def issue(sheet: str, item: str, value: Any, status: str, blocking: bool, reason: str, action: str) -> dict[str, Any]:
    return {
        "sheet": sheet,
        "item": item,
        "value": as_text(value),
        "status": status or "缺失",
        "blocking": blocking,
        "reason": reason,
        "action": action,
    }


def validate_fields(records: dict[str, dict[str, Any]], specs: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    values = {key: record["value"] for key, record in records.items()}
    for field_id, record in records.items():
        spec = specs.get(field_id, {})
        if not condition_active(record["applicability"], values):
            continue
        status = record["status"]
        value = record["value"]
        critical = record["critical"]
        allow_na = bool(spec.get("allow_na"))
        valid = status == "已确认" and bool(value)
        if status == "不适用" and allow_na:
            valid = True
        if critical and not valid:
            reason = "关键字段必须有值并由作者确认。"
            if status == "不适用" and not allow_na:
                reason = "该关键字段不能以不适用完成。"
            issues.append(issue(record["sheet"], record["zh"], value, status, True, reason, "核对真实信息后标为已确认；仅在字段允许时选择不适用。"))
        elif not critical and status in {"缺失", "抽取待确认"}:
            issues.append(issue(record["sheet"], record["zh"], value, status, False, "可选字段尚未确认。", "当前期刊要求时补充；已有抽取值时请核对。"))
    return issues


def validate_entities(wb) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    authors = rows_as_dicts(wb[SHEET_AUTHORS])
    affiliations = rows_as_dicts(wb[SHEET_AFFILIATIONS])
    contributions = rows_as_dicts(wb[SHEET_CONTRIBUTIONS])

    if not authors:
        issues.append(issue(SHEET_AUTHORS, "作者清单", "", "缺失", True, "至少需要一位作者。", "填写作者顺序、显示名、单位和角色。"))
        return issues

    author_ids: set[str] = set()
    orders: list[str] = []
    corresponding_count = 0
    submitter_count = 0
    for position, author in enumerate(authors, start=2):
        author_id = as_text(author.get("作者ID")) or f"第{position - 1}位作者"
        author_ids.add(author_id)
        order = as_text(author.get("作者顺序"))
        if order:
            orders.append(order)
        required = ["作者ID", "作者顺序", "投稿显示名", "单位ID"]
        missing = [name for name in required if not as_text(author.get(name))]
        status = as_text(author.get("记录状态"))
        if missing or status != "已确认":
            issues.append(issue(
                SHEET_AUTHORS, author_id, author.get("投稿显示名"), status, True,
                ("缺少：" + "、".join(missing) + "；" if missing else "") + "作者记录尚未确认。",
                "核对姓名、顺序、单位映射和作者角色后将记录状态设为已确认。",
            ))
        if normalize_yes(author.get("通讯作者")) == "是":
            corresponding_count += 1
            if not as_text(author.get("邮箱")):
                issues.append(issue(SHEET_AUTHORS, f"{author_id} 通讯邮箱", "", status, True, "通讯作者缺少邮箱。", "填写作者明确提供的邮箱。"))
        if normalize_yes(author.get("提交人")) == "是":
            submitter_count += 1

    if len(orders) != len(set(orders)):
        issues.append(issue(SHEET_AUTHORS, "作者顺序", ", ".join(orders), "抽取待确认", True, "作者顺序存在重复。", "确认唯一作者顺序。"))
    if corresponding_count == 0:
        issues.append(issue(SHEET_AUTHORS, "通讯作者", "", "缺失", True, "尚未明确通讯作者。", "由作者明确指定至少一位通讯作者。"))
    if submitter_count == 0:
        issues.append(issue(SHEET_AUTHORS, "提交人", "", "缺失", True, "尚未明确投稿系统提交人。", "由作者明确指定提交人。"))

    affiliation_ids: set[str] = set()
    for affiliation in affiliations:
        aff_id = as_text(affiliation.get("单位ID"))
        if aff_id:
            affiliation_ids.add(aff_id)
        missing = [name for name in ["单位ID", "机构", "国家/地区"] if not as_text(affiliation.get(name))]
        status = as_text(affiliation.get("记录状态"))
        if missing or status != "已确认":
            issues.append(issue(
                SHEET_AFFILIATIONS, aff_id or "未编号单位", affiliation.get("机构"), status, True,
                ("缺少：" + "、".join(missing) + "；" if missing else "") + "单位记录尚未确认。",
                "核对单位全名、地址和作者映射后将记录状态设为已确认。",
            ))
    if not affiliations:
        issues.append(issue(SHEET_AFFILIATIONS, "作者单位", "", "缺失", True, "作者单位清单为空。", "填写单位并与作者ID关联。"))
    for author in authors:
        author_id = as_text(author.get("作者ID"))
        refs = [item.strip() for item in as_text(author.get("单位ID")).replace("；", ",").replace(";", ",").split(",") if item.strip()]
        missing_refs = [ref for ref in refs if ref not in affiliation_ids]
        if missing_refs:
            issues.append(issue(SHEET_AFFILIATIONS, f"{author_id} 单位映射", ", ".join(missing_refs), "缺失", True, "作者引用了不存在的单位ID。", "补充单位记录或修正作者的单位ID。"))

    contribution_by_author = {as_text(row.get("作者ID")): row for row in contributions if as_text(row.get("作者ID"))}
    for author_id in sorted(author_ids):
        row = contribution_by_author.get(author_id)
        if row is None:
            issues.append(issue(SHEET_CONTRIBUTIONS, author_id, "", "缺失", True, "缺少该作者的贡献与批准记录。", "填写 CRediT 角色和作者确认。"))
            continue
        status = as_text(row.get("记录状态"))
        confirmations = ["作者资格", "批准最终稿", "同意投稿", "责任确认"]
        invalid_confirmations = [name for name in confirmations if normalize_yes(row.get(name)) != "是"]
        unset_roles = [role for role in CREDIT_ROLES if normalize_yes(row.get(role)) not in {"是", "否", "不适用"}]
        selected_roles = [role for role in CREDIT_ROLES if normalize_yes(row.get(role)) == "是"]
        reasons: list[str] = []
        if invalid_confirmations:
            reasons.append("未确认：" + "、".join(invalid_confirmations))
        if unset_roles:
            reasons.append("CRediT 未逐项填写")
        if not selected_roles:
            reasons.append("未选择任何贡献角色")
        if status != "已确认":
            reasons.append("记录尚未确认")
        if reasons:
            issues.append(issue(SHEET_CONTRIBUTIONS, author_id, "; ".join(selected_roles), status, True, "；".join(reasons) + "。", "由作者核对贡献、批准和责任确认。"))

    known_contributors = set(contribution_by_author)
    for extra in sorted(known_contributors - author_ids):
        issues.append(issue(SHEET_CONTRIBUTIONS, extra, "", "抽取待确认", True, "贡献表中的作者ID不在作者清单中。", "修正作者ID或补充作者记录。"))

    for reviewer in rows_as_dicts(wb[SHEET_REVIEWERS]):
        record_id = as_text(reviewer.get("记录ID")) or "未编号记录"
        status = as_text(reviewer.get("记录状态"))
        if status != "已确认" or not as_text(reviewer.get("姓名")):
            issues.append(issue(SHEET_REVIEWERS, record_id, reviewer.get("姓名"), status, False, "编辑/审稿人记录尚未确认或不完整。", "仅在期刊要求时补齐姓名、机构、邮箱、理由和关系披露。"))

    for file_row in rows_as_dicts(wb[SHEET_FILES]):
        file_id = as_text(file_row.get("文件ID")) or "未编号文件"
        if normalize_yes(file_row.get("是否必需")) == "是" and as_text(file_row.get("准备状态")) != "已准备":
            issues.append(issue(SHEET_FILES, file_id, file_row.get("文件名"), as_text(file_row.get("准备状态")), False, "标记为必需的文件尚未准备。", "完成文件并更新准备状态；具体要求以目标期刊为准。"))
    return issues


def write_issues(wb, issues: list[dict[str, Any]]) -> None:
    ws = wb[SHEET_ISSUES]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for index, item in enumerate(issues, start=1):
        ws.append([
            index, item["sheet"], item["item"], item["value"], item["status"],
            "是" if item["blocking"] else "否", item["reason"], item["action"],
        ])
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 2)}"


def field_visibility(sheet: str, field_name: str, fallback: str) -> str:
    public_author_fields = {
        "作者顺序", "名/Given name", "中间名/Middle name", "姓/Family name", "投稿显示名",
        "学位", "职位", "单位ID", "通讯作者", "共同第一作者", "共同资深作者", "责任作者/保证人",
    }
    if sheet == SHEET_AUTHORS:
        return "论文公开" if field_name in public_author_fields else "仅投稿系统"
    if sheet in {SHEET_AFFILIATIONS, SHEET_CONTRIBUTIONS}:
        return "论文公开"
    return fallback or "内部使用"


def sync_field_metadata(wb) -> None:
    if SHEET_METADATA not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_METADATA)
        append_headers(ws, METADATA_HEADERS)
        style_sheet(ws, {"A": 20, "B": 18, "C": 30, "D": 45, "E": 16, "F": 24, "G": 16, "H": 42, "I": 16})
    ws = wb[SHEET_METADATA]
    ws.sheet_state = "visible"
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    records, _ = get_field_records(wb)
    for record in records.values():
        ws.append([
            record["sheet"], record["field_id"], record["zh"], record["value"], record["status"],
            record["source"], record["visibility"], record["note"], record["confirmed_date"],
        ])

    entity_specs = {
        SHEET_AUTHORS: "作者ID",
        SHEET_AFFILIATIONS: "单位ID",
        SHEET_CONTRIBUTIONS: "作者ID",
        SHEET_REVIEWERS: "记录ID",
        SHEET_FILES: "文件ID",
    }
    metadata_fields = {"记录状态", "来源", "可见性", "备注", "最后确认日期"}
    for sheet, id_field in entity_specs.items():
        for position, record in enumerate(rows_as_dicts(wb[sheet]), start=1):
            object_id = as_text(record.get(id_field)) or f"row-{position}"
            status = as_text(record.get("记录状态")) or "缺失"
            source = as_text(record.get("来源"))
            visibility = as_text(record.get("可见性")) or "内部使用"
            note = as_text(record.get("备注"))
            confirmed_date = as_text(record.get("最后确认日期"))
            for name, value in record.items():
                if name in metadata_fields or name == id_field:
                    continue
                ws.append([
                    sheet, object_id, name, as_text(value), status, source,
                    field_visibility(sheet, name, visibility), note, confirmed_date,
                ])
    ws.auto_filter.ref = f"A1:I{max(ws.max_row, 2)}"
    ws.sheet_state = "hidden"


def update_status_sheet(wb, blocking: int, nonblocking: int) -> str:
    status = "最终版" if blocking == 0 else "草稿版"
    ws = wb[SHEET_STATUS]
    values = {
        "模板版本": VERSION,
        "生成日期": today_text(),
        "文档状态": status,
        "阻塞项数量": blocking,
        "非阻塞待办数量": nonblocking,
    }
    for row in range(2, ws.max_row + 1):
        key = as_text(ws.cell(row, 1).value)
        if key in values:
            ws.cell(row, 2).value = values[key]
    return status


def validate_workbook(path: Path, *, save: bool = True) -> dict[str, Any]:
    wb = load_workbook(path)
    required = [
        SHEET_STATUS, SHEET_MANUSCRIPT, SHEET_AUTHORS, SHEET_AFFILIATIONS,
        SHEET_CONTRIBUTIONS, SHEET_DECLARATIONS, SHEET_ETHICS, SHEET_DATA,
        SHEET_REVIEWERS, SHEET_FILES, SHEET_ISSUES,
    ]
    missing_sheets = [name for name in required if name not in wb.sheetnames]
    if missing_sheets:
        raise ValueError("Workbook is missing required sheets: " + ", ".join(missing_sheets))
    records, specs = get_field_records(wb)
    issues = validate_fields(records, specs) + validate_entities(wb)
    blocking = sum(1 for item in issues if item["blocking"])
    nonblocking = len(issues) - blocking
    write_issues(wb, issues)
    document_status = update_status_sheet(wb, blocking, nonblocking)
    sync_field_metadata(wb)
    if save:
        wb.save(path)
    return {
        "workbook": str(path.resolve()),
        "document_status": document_status,
        "blocking_count": blocking,
        "nonblocking_count": nonblocking,
        "issue_count": len(issues),
        "issues": issues,
    }


def set_doc_defaults(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)


def add_table(document: Document, headers: list[str], rows: Iterable[Iterable[Any]], *, font_size: float = 8.5) -> None:
    rows_list = [list(row) for row in rows]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(font_size)
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(table_header)
    for values in rows_list:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = as_text(value)
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)


def field_rows_for_word(records: dict[str, dict[str, Any]], sheet: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for spec in FIELD_SPECS[sheet]:
        record = records.get(spec["id"])
        if not record:
            continue
        value = record["value"]
        status = record["status"]
        source = record["source"]
        condition = record.get("applicability", "全部")
        if condition != "全部" and "=" in condition:
            controller_id, expected = condition.split("=", 1)
            controller = records.get(controller_id.strip(), {})
            allowed = {part.strip() for part in expected.split("|")}
            controller_value = as_text(controller.get("value"))
            controller_status = as_text(controller.get("status"))
            if controller_status == "已确认" and controller_value not in allowed:
                value = "—"
                status = "不适用（由控制项判定）"
                source = controller.get("source") or "控制项"
            elif controller_status != "已确认":
                status = "等待控制项确认"
        rows.append([record["zh"], record["en"], value, status, source])
    return rows


def status_value(wb, key: str) -> str:
    ws = wb[SHEET_STATUS]
    for row in range(2, ws.max_row + 1):
        if as_text(ws.cell(row, 1).value) == key:
            return as_text(ws.cell(row, 2).value)
    return ""


def render_document(workbook_path: Path, output_path: Path, *, force: bool = False) -> dict[str, Any]:
    if output_path.exists() and not force:
        raise FileExistsError(f"Output exists and is preserved: {output_path}")
    summary = validate_workbook(workbook_path, save=True)
    wb = load_workbook(workbook_path, data_only=False)
    records, _ = get_field_records(wb)
    title = records.get("manuscript_title", {}).get("value") or "未命名稿件"

    document = Document()
    set_doc_defaults(document)
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("论文投稿信息参照文档\nSubmission Information Reference")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(title).bold = True
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"版本：{VERSION}    生成日期：{today_text()}    状态：{summary['document_status']}")
    warning = document.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.add_run("隐私提示：本文档可能含个人及仅投稿系统使用的信息，请限制访问。\n").bold = True
    warning.add_run("不应在本文档中保存密码、验证码、登录令牌、银行卡或支付凭据。")
    document.add_page_break()

    document.add_heading("1. 投稿准备摘要 / Submission readiness", level=1)
    add_table(document, ["项目", "当前状态"], [
        ["文档状态", summary["document_status"]],
        ["阻塞最终版的事项", summary["blocking_count"]],
        ["非阻塞待办", summary["nonblocking_count"]],
        ["资料源", "Excel 投稿信息总表"],
        ["适用范围", "跨期刊首次投稿通用资料；不替代目标期刊现行指南"],
    ], font_size=9)
    if summary["document_status"] != "最终版":
        document.add_paragraph("当前为草稿版。关键事实尚未全部确认，不应将本文档视为投稿就绪证明。")

    document.add_heading("2. 稿件基本信息 / Manuscript information", level=1)
    add_table(document, ["中文字段", "English field", "值", "状态", "来源"], field_rows_for_word(records, SHEET_MANUSCRIPT))

    authors = rows_as_dicts(wb[SHEET_AUTHORS])
    document.add_heading("3. 作者与投稿角色 / Authors and submission roles", level=1)
    if authors:
        add_table(document, ["顺序", "作者", "单位ID", "邮箱", "ORCID", "通讯", "提交人", "保证人", "状态"], [
            [row.get("作者顺序"), row.get("投稿显示名"), row.get("单位ID"), row.get("邮箱"), row.get("ORCID"),
             row.get("通讯作者"), row.get("提交人"), row.get("责任作者/保证人"), row.get("记录状态")]
            for row in authors
        ], font_size=7.5)
        document.add_heading("仅投稿系统使用的作者资料", level=2)
        add_table(document, ["作者", "学位", "职位", "称谓", "性别/系统选项", "国家/地区", "电话"], [
            [row.get("投稿显示名"), row.get("学位"), row.get("职位"), row.get("称谓"), row.get("性别或系统选项"),
             row.get("国家/地区"), row.get("电话")]
            for row in authors
        ], font_size=7.5)
    else:
        document.add_paragraph("尚未填写作者信息。")

    document.add_heading("4. 作者单位 / Affiliations", level=1)
    affiliations = rows_as_dicts(wb[SHEET_AFFILIATIONS])
    if affiliations:
        add_table(document, ["单位ID", "机构", "科室/部门", "地址", "国家/地区", "关联作者ID", "状态"], [
            [row.get("单位ID"), row.get("机构"), row.get("科室/部门"),
             ", ".join(filter(None, [as_text(row.get("城市")), as_text(row.get("省/州")), as_text(row.get("邮编"))])),
             row.get("国家/地区"), row.get("关联作者ID"), row.get("记录状态")]
            for row in affiliations
        ], font_size=8)
    else:
        document.add_paragraph("尚未填写作者单位。")

    document.add_heading("5. 作者贡献与批准 / Contributions and approvals", level=1)
    contributions = rows_as_dicts(wb[SHEET_CONTRIBUTIONS])
    if contributions:
        contribution_rows = []
        for row in contributions:
            selected = [role for role in CREDIT_ROLES if normalize_yes(row.get(role)) == "是"]
            contribution_rows.append([
                row.get("作者ID"), "; ".join(selected), row.get("作者资格"), row.get("批准最终稿"),
                row.get("同意投稿"), row.get("责任确认"), row.get("记录状态"),
            ])
        add_table(document, ["作者ID", "CRediT roles", "作者资格", "批准终稿", "同意投稿", "责任确认", "状态"], contribution_rows, font_size=7.5)
    else:
        document.add_paragraph("尚未填写贡献与批准信息。")

    document.add_heading("6. 投稿声明 / Declarations", level=1)
    add_table(document, ["中文字段", "English field", "值", "状态", "来源"], field_rows_for_word(records, SHEET_DECLARATIONS), font_size=8)
    document.add_heading("已确认、可直接粘贴的英文声明", level=2)
    confirmed_statements = []
    for spec in FIELD_SPECS[SHEET_DECLARATIONS]:
        record = records.get(spec["id"], {})
        if spec["id"].endswith("_en") and record.get("status") == "已确认" and record.get("value"):
            confirmed_statements.append([record.get("en"), record.get("value")])
    if confirmed_statements:
        add_table(document, ["Statement", "Confirmed text"], confirmed_statements, font_size=9)
    else:
        document.add_paragraph("当前没有已确认、可直接粘贴的英文声明。待确认内容未被转换为确定声明。")

    document.add_heading("7. 伦理与注册 / Ethics and registration", level=1)
    add_table(document, ["中文字段", "English field", "值", "状态", "来源"], field_rows_for_word(records, SHEET_ETHICS), font_size=8)

    document.add_heading("8. 数据、代码与材料 / Data, code, and materials", level=1)
    add_table(document, ["中文字段", "English field", "值", "状态", "来源"], field_rows_for_word(records, SHEET_DATA), font_size=8)

    document.add_heading("9. 编辑与审稿人 / Editors and reviewers", level=1)
    reviewers = rows_as_dicts(wb[SHEET_REVIEWERS])
    if reviewers:
        add_table(document, ["类型", "姓名", "机构", "邮箱", "理由", "关系披露", "状态"], [
            [row.get("类型"), row.get("姓名"), row.get("机构"), row.get("邮箱"), row.get("理由"), row.get("关系披露"), row.get("记录状态")]
            for row in reviewers
        ], font_size=7.5)
    else:
        document.add_paragraph("尚未记录推荐/回避审稿人或编辑；仅在目标期刊要求时补充。")

    document.add_heading("10. 上传文件清单 / Upload inventory", level=1)
    files = rows_as_dicts(wb[SHEET_FILES])
    if files:
        add_table(document, ["文件类型", "文件名", "版本", "匿名", "必需", "准备状态", "记录状态"], [
            [row.get("文件类型"), row.get("文件名"), row.get("版本"), row.get("是否匿名"), row.get("是否必需"), row.get("准备状态"), row.get("记录状态")]
            for row in files
        ], font_size=8)
    else:
        document.add_paragraph("尚未建立上传文件清单。")

    document.add_heading("11. 缺项与待确认 / Missing and pending items", level=1)
    issues = rows_as_dicts(wb[SHEET_ISSUES])
    if issues:
        add_table(document, ["工作表", "对象/字段", "状态", "阻塞", "原因", "下一步"], [
            [row.get("工作表"), row.get("对象/字段"), row.get("状态"), row.get("是否阻塞最终版"), row.get("原因"), row.get("下一步")]
            for row in issues
        ], font_size=7.5)
    else:
        document.add_paragraph("没有阻塞项或待确认项。")

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "Submission information reference — private working document"
        for run in footer.runs:
            run.font.size = Pt(8)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {**summary, "document": str(output_path.resolve())}


def init_command(args: argparse.Namespace) -> dict[str, Any]:
    output = args.output.resolve()
    if output.exists() and not args.force:
        raise FileExistsError(f"Output exists and is preserved: {output}")
    wb = create_workbook()
    warnings: list[str] = []
    if args.seed_json:
        with args.seed_json.open("r", encoding="utf-8-sig") as handle:
            seed = json.load(handle)
        if not isinstance(seed, dict):
            raise ValueError("Seed JSON root must be an object.")
        warnings = apply_seed(wb, seed)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    summary = validate_workbook(output, save=True)
    return {**summary, "warnings": warnings}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    init_parser = subparsers.add_parser("init", help="Create a new master workbook, optionally prefilled from temporary JSON.")
    init_parser.add_argument("--output", type=Path, required=True)
    init_parser.add_argument("--seed-json", type=Path)
    init_parser.add_argument("--force", action="store_true")

    validate_parser = subparsers.add_parser("validate", help="Refresh readiness and missing-item sheets.")
    validate_parser.add_argument("--workbook", type=Path, required=True)

    render_parser = subparsers.add_parser("render", help="Validate the workbook and render the bilingual Word reference.")
    render_parser.add_argument("--workbook", type=Path, required=True)
    render_parser.add_argument("--output", type=Path, required=True)
    render_parser.add_argument("--force", action="store_true")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        if args.command == "init":
            result = init_command(args)
        elif args.command == "validate":
            result = validate_workbook(args.workbook.resolve(), save=True)
        elif args.command == "render":
            result = render_document(args.workbook.resolve(), args.output.resolve(), force=args.force)
        else:
            parser.error("Unknown command")
            return 2
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1
    print(json.dumps({"ok": True, **result}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
