"""Create a basic editable PPTX and speaker notes from supplied JSON."""
import argparse
import io
import json
from pathlib import Path


def build(source, output):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from PIL import Image

    data = json.loads(source.read_text(encoding='utf-8-sig'))
    slides = data['slides']
    if not isinstance(slides, list) or not slides:
        raise ValueError('slides must be a non-empty array')
    if output.exists():
        raise FileExistsError(output)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = data.get('title', '')
    prs.core_properties.author = data.get('author', '')
    for index, spec in enumerate(slides, 1):
        if not str(spec.get('title', '')).strip():
            raise ValueError(f'Slide {index} needs a title')
        bullets = spec.get('bullets', [])
        if not isinstance(bullets, list) or any(not isinstance(x, str) for x in bullets):
            raise ValueError('bullets must contain strings')
        # This basic layout is intentionally bounded; split dense material into more slides.
        if len(bullets) > 6 or any(len(x) > 120 for x in bullets) or len(spec['title']) > 70:
            raise ValueError(f'Slide {index} is too dense for the basic layout; shorten or split it')
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = RGBColor.from_string('F4F7FB')
        def text_box(x, y, width, height, text, size, color='17324D', bold=False):
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            frame = shape.text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = Inches(0.02)
            para = frame.paragraphs[0]
            para.text = text
            para.font.name = 'Microsoft YaHei'
            para.font.size = Pt(size)
            para.font.bold = bold
            para.font.color.rgb = RGBColor.from_string(color)
            return shape
        text_box(0.7, 0.45, 11.95, 1.0, spec['title'], 30, bold=True)
        image_name = spec.get('image')
        content_width = 6.0 if image_name else 11.8
        text_box(0.8, 1.65, content_width, 4.85, '\n\n'.join('• ' + b for b in bullets), 21)
        if image_name:
            path = Path(image_name)
            if not path.is_absolute():
                path = source.parent / path
            with Image.open(path) as im:
                width, height = im.size
            factor = min(5.0 / width, 4.85 / height)
            w, h = width * factor, height * factor
            slide.shapes.add_picture(str(path), Inches(7.35 + (5 - w) / 2), Inches(1.65 + (4.85 - h) / 2), width=Inches(w), height=Inches(h))
        text_box(0.8, 7.00, 11.7, 0.28, f'{data.get("title", "")}  ·  {index}/{len(slides)}', 10, '62748A')
        slide.notes_slide.notes_text_frame.text = str(spec.get('notes', ''))
    buffer = io.BytesIO()
    prs.save(buffer)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as handle:
        handle.write(buffer.getvalue())
    print(f'Created {len(slides)} slides; text and notes saved, visual layout still requires rendering: {output.resolve()}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.input, args.output)


if __name__ == '__main__':
    main()
